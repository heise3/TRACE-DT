from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
p=R/'deliverables/upload/Manuscript.docx';d=Document(p)
for pp in d.paragraphs:
 if pp.style.name not in ['Title']:pp.paragraph_format.line_spacing=2.0
for sec in d.sections:
 if not sec._sectPr.xpath('./w:lnNumType'):
  tag=OxmlElement('w:lnNumType');tag.set(qn('w:countBy'),'1');tag.set(qn('w:restart'),'continuous');tag.set(qn('w:distance'),'240');sec._sectPr.append(tag)
for t in d.tables:
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
for dd in [d]:
 for sec in dd.sections:
  for part in [sec.header,sec.footer,sec.even_page_header,sec.even_page_footer,sec.first_page_header,sec.first_page_footer]:
   for pp in part.paragraphs:pp._p.get_or_add_pPr().append(OxmlElement('w:suppressLineNumbers'))
for pp in d.paragraphs:
 if pp.text.startswith('Table 1. Methods compared'):pp.paragraph_format.keep_with_next=True
for t in d.tables:
 for i,row in enumerate(t.rows):
  for cell in row.cells:
   for pp in cell.paragraphs:pp.paragraph_format.keep_with_next=(i<len(t.rows)-1)
d.save(p)
p=R/'deliverables/editable/Supplementary_Methods.docx';s=Document(p)
for sec in s.sections:
 for part in [sec.header,sec.footer,sec.even_page_header,sec.even_page_footer,sec.first_page_header,sec.first_page_footer]:
  for pp in part.paragraphs:pp._p.get_or_add_pPr().append(OxmlElement('w:suppressLineNumbers'))
s.save(p)
