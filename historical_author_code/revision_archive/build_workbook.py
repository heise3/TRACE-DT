from pathlib import Path
import csv,json,re,hashlib,shutil
from openpyxl import Workbook,load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill
R=Path(__file__).resolve().parents[1];D=R/'deliverables';UP=D/'upload';A=D/'revision_archive';A.mkdir(exist_ok=True)
source=R/'working/frozen/tables'
files=list(source.glob('Table_S*.tsv'))+list((source/'R39_S40_S45').glob('*.tsv'))+list((source/'R39_S40_S45').glob('S*.json'))+list((R/'analysis').glob('Table_S*.tsv'))
files=sorted(files,key=lambda p:(int(re.search(r'S(\d+)',p.name).group(1)) if re.search(r'S(\d+)',p.name) else 100,p.name))
wb=Workbook(write_only=True);index=wb.create_sheet('Index');rows=[];taken=set();checks=[]
def digest_rows(items):
 h=hashlib.sha256()
 for row in items:h.update(json.dumps(row,ensure_ascii=False,separators=(',',':')).encode());h.update(b'\n')
 return h.hexdigest()
for p in files:
 m=re.search(r'S(\d+)',p.name);key=m.group(0) if m else 'Metadata'
 n=1;sheet=key
 while sheet in taken:n+=1;sheet=f'{key}_{n}'
 taken.add(sheet);ws=wb.create_sheet(sheet);ws.freeze_panes='A2'
 if p.suffix=='.json':
  obj=json.loads(p.read_text());data=[['json_path','value_json']]
  def flatten(x,path='$'):
   if isinstance(x,dict):
    for k,v in x.items():yield from flatten(v,path+'.'+k)
   elif isinstance(x,list):
    for i,v in enumerate(x):yield from flatten(v,f'{path}[{i}]')
   else:yield [path,json.dumps(x,ensure_ascii=False)]
  data+=list(flatten(obj));it=iter(data)
 else:it=csv.reader(p.open(newline='',encoding='utf-8-sig'),delimiter='\t')
 h=hashlib.sha256();count=0;width=None
 for rr in it:
  if width is None:width=len(rr)
  assert len(rr)==width,(p,count,len(rr),width)
  assert max(map(len,rr),default=0)<=32767,(p,count,'cell too long')
  h.update(json.dumps(rr,ensure_ascii=False,separators=(',',':')).encode());h.update(b'\n')
  cells=[]
  for v in rr:
   c=WriteOnlyCell(ws,value=v);c.data_type='s' # preserve IDs, zeros, null strings and formula-like text exactly
   if count==0:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='197C80')
   cells.append(c)
  ws.append(cells);count+=1
 ws.auto_filter.ref=f'A1:{__import__("openpyxl").utils.get_column_letter(width)}{count}'
 rows.append([sheet,p.name,str(count-1),hashlib.sha256(p.read_bytes()).hexdigest(),'All fields stored as source text; empty strings remain empty.'])
 checks.append({'sheet':sheet,'file':p.name,'rows_with_header':count,'columns':width,'content_sha256':h.hexdigest()})
 if p.parent==R/'analysis':shutil.copy2(p,A/p.name)
index.append(['sheet','source_file','data_rows','source_sha256','representation'])
for row in rows:index.append(row)
path=UP/'Supplementary_Tables.xlsx';wb.save(path)
# Round-trip verify every table cell; read-only iteration avoids Excel type coercion.
rd=load_workbook(path,read_only=True,data_only=False)
for check in checks:
 h=hashlib.sha256();n=0
 for row in rd[check['sheet']].iter_rows(values_only=True):
  rr=['' if x is None else str(x) for x in row]
  h.update(json.dumps(rr,ensure_ascii=False,separators=(',',':')).encode());h.update(b'\n');n+=1
 assert n==check['rows_with_header'] and h.hexdigest()==check['content_sha256'],check
rd.close()
(R/'deliverables/checks/Workbook_QA.json').write_text(json.dumps({'all_cells_roundtrip':'PASS','sheets':len(checks)+1,'total_data_rows':sum(c['rows_with_header']-1 for c in checks),'tables':checks},indent=2))
print('PASS',len(checks)+1,'sheets',sum(c['rows_with_header']-1 for c in checks),'data rows',path.stat().st_size)
